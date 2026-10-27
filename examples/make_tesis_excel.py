"""Genera el Excel de ejemplo de la columna de campo de la tesis.

Uso:  python examples/make_tesis_excel.py

Columna de afloramiento (alturas desde la base, unidades listadas desde la base):
esquisto (30 m), pizarra (15 m) y andesita (60 m). Los contactos son inferidos.
Las tres vetas de cuarzo van en la misma hoja con `kind` = vein.

IMPORTANTE: las alturas de las unidades y las posiciones de las vetas son
valores de EJEMPLO. Del campo solo se tomó el espesor de las vetas (60, 27 y
15 cm). Reemplaza las posiciones por las medidas reales.

`top_contact` describe el contacto en el TECHO de la unidad (con la unidad que
tiene encima); el techo de la andesita no tiene unidad encima, así que se omite.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = ["unit_id", "kind", "base_height_m", "top_height_m", "thickness_m",
           "lithology", "top_contact", "observations"]

N = None
ROWS = [
    ("T1", "unit", 0.0, 30.0, N, "esquisto", "inferido", "Altura de ejemplo (30 m)"),
    ("T2", "unit", 30.0, 45.0, N, "pizarra", "inferido", "Pizarra calcárea. Altura de ejemplo (15 m)"),
    ("T3", "unit", 45.0, 105.0, N, "andesita", N, "Altura de ejemplo (60 m)"),
    ("V1", "vein", 33.0, N, 0.60, "veta_cuarzo", N, "Veta de espesor máximo (60 cm). Posición ILUSTRATIVA"),
    ("V2", "vein", 38.0, N, 0.15, "veta_cuarzo", N, "Veta de 15 cm. Posición ILUSTRATIVA"),
    ("V3", "vein", 42.5, N, 0.27, "veta_cuarzo", N, "Veta de 27 cm. Posición ILUSTRATIVA"),
]


def build(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Columna_Tesis"
    ws.append(HEADERS)
    for row in ROWS:
        ws.append(list(row))

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", start_color="4C6B52", end_color="4C6B52")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.border = border
    for i, header in enumerate(HEADERS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 55 if header == "observations" else 16
    ws.freeze_panes = "A2"
    wb.save(path)


if __name__ == "__main__":
    out = Path(__file__).parent / "ejemplo_tesis.xlsx"
    build(out)
    print(f"Generado: {out}")
