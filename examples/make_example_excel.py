"""Genera el Excel de ejemplo (secuencia volcano-sedimentaria con alteración).

Uso:  python examples/make_example_excel.py

Las columnas `color` y `pattern` se dejan vacías a propósito en casi todas
las filas: GeoColumn toma el color y el patrón por defecto de la litología
(ver geocolumn/styles/lithologies.py). Solo U2 sobrescribe el color, para
mostrar que lo que se escribe en la unidad tiene prioridad sobre el catálogo.

`top_contact` describe el contacto en el TECHO de la unidad, es decir, con la
unidad que tiene encima. Por eso la falla que limita la base de U5 está
escrita en U6 (la unidad de abajo).
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = [
    "unit_id", "top_depth_m", "base_depth_m", "thickness_m",
    "lithology", "color", "pattern", "grain_size",
    "sedimentary_structures", "fossils", "samples",
    "hydrothermal_alteration", "mineralization",
    "top_contact",
    "age_unit_name", "age_min_ma", "age_max_ma", "age_period", "age_epoch",
    "photos", "observations",
]

N = None
ROWS = [
    ["U1", 0, 15, N, "andesita_lava", N, N, N, "flow_banding", N, "M-01;M-02",
     "propilitica_debil", "trazas_pirita_diseminada", "sharp",
     "Fm. Volcánica Superior", 9.0, 10.5, "Neógeno", "Mioceno", N,
     "Lavas frescas a débilmente alteradas; clorita y epidota en fracturas"],
    ["U2", 15, 28, N, "brecha_volcanica", "gris_oscuro", N, "bloques a lapilli", "clast_supported", N, "M-03",
     "propilitica", "pirita_diseminada_2pct", "gradual",
     "Fm. Volcánica Superior", 9.5, 11.0, "Neógeno", "Mioceno", N,
     "Brecha volcánica con clastos andesíticos angulosos (color sobrescrito a propósito)"],
    ["U3", 28, 40, N, "arenisca_tobacea", N, N, "arena media", "cross_bedding;ripple_marks",
     "restos_vegetales_carbonizados", "M-04;M-05", N, N, "sharp",
     "Fm. Sedimentaria Intermedia", 10.0, 11.5, "Neógeno", "Mioceno", N,
     "Nivel sedimentario intercalado con retrabajo de material piroclástico"],
    ["U4", 40, 55, N, "andesita_lava", N, N, N, "vesicular_zone", N, "M-06",
     "argilica_moderada", "pirita_diseminada_3-5pct;trazas_calcopirita", "sharp",
     "Fm. Volcánica Inferior", 11.0, 12.5, "Neógeno", "Mioceno", N,
     "Alteración argílica sobreimpuesta; illita y caolinita en la matriz"],
    ["U5", 55, 70, N, "brecha_hidrotermal", N, N, "matriz fina", "vein_stockwork", N, "M-07;M-08;M-09",
     "silicificacion_intensa;argilica_avanzada", "calcopirita_pirita_5-8pct;vetillas_cuarzo_sulfuros", "sharp",
     "Fm. Volcánica Inferior", 11.5, 13.0, "Neógeno", "Mioceno", N,
     "Brecha hidrotermal mineralizada con stockwork de cuarzo-sulfuros; falla en la base"],
    ["U6", 70, 85, N, "arenisca_volcaniclastica", N, N, "arena gruesa", "planar_bedding", N, "M-10",
     "propilitica_debil", "trazas_pirita", "fault",
     "Fm. Sedimentaria Basal", 12.0, 13.5, "Neógeno", "Mioceno", N,
     "Retrabajo sedimentario de material volcánico; alteración débil"],
    ["U7", 85, 100, N, "andesita_lava", N, N, N, "columnar_jointing", N, "M-11",
     "propilitica", N, "sharp",
     "Fm. Volcánica Basal", 13.0, 14.5, "Neógeno", "Mioceno", N,
     "Lavas con clorita, epidota y calcita; sin mineralización significativa"],
    ["U8", 100, 118, N, "domo_dacitico", N, N, N, "flow_banding", N, "M-12;M-13",
     "propilitica_debil", "trazas_pirita", "sharp",
     "Complejo Subvolcánico", 14.0, 15.5, "Neógeno", "Mioceno", N,
     "Domo dacítico subvolcánico, probable fuente de calor del sistema; base del sondaje"],
]


def build(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Columna_Ejemplo"
    ws.append(HEADERS)
    for row in ROWS:
        ws.append(row)

    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for cell in ws[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", start_color="4C6B52", end_color="4C6B52")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.border = border
    for i, header in enumerate(HEADERS, start=1):
        width = 45 if header == "observations" else 22 if "age_unit" in header or header in (
            "lithology", "mineralization", "hydrothermal_alteration") else 14
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    wb.save(path)


if __name__ == "__main__":
    out = Path(__file__).parent / "ejemplo_volcanosedimentario.xlsx"
    build(out)
    print(f"Generado: {out}")
